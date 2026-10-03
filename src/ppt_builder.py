"""Automated PowerPoint report generation with python-pptx."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from . import analytics, config

NAVY = RGBColor(0x1F, 0x4E, 0x79)
ORANGE = RGBColor(0xF1, 0x8F, 0x01)
GREY = RGBColor(0x55, 0x55, 0x55)
LIGHT = RGBColor(0xEA, 0xF1, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SW, SH = Inches(13.333), Inches(7.5)  # 16:9


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _text(slide, text, left, top, width, height, size=14, bold=False, color=GREY,
          align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.size, p.font.bold, p.font.color.rgb = Pt(size), bold, color
    return tb


def _header(slide, title, subtitle=None):
    bar = slide.shapes.add_shape(1, 0, 0, SW, Inches(1.0))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    _text(slide, title, Inches(0.5), Inches(0.18), Inches(12), Inches(0.7), 28, True, WHITE)
    if subtitle:
        _text(slide, subtitle, Inches(0.5), Inches(1.08), Inches(12), Inches(0.4), 13, False, GREY)


def _bullets(slide, lines, left, top, width, height, size=16):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "\u2022  " + line
        p.font.size, p.font.color.rgb = Pt(size), GREY
        p.space_after = Pt(10)


def _picture_fit(slide, path, left, top, max_w, max_h, center=True):
    """Insert a picture scaled to fit inside the box without distortion."""
    from PIL import Image
    with Image.open(path) as im:
        w, h = im.size
    scale = min(max_w / w, max_h / h)
    pw, ph = int(w * scale), int(h * scale)
    x = left + (max_w - pw) // 2 if center else left
    slide.shapes.add_picture(str(path), x, top, pw, ph)


def _footer(slide, n):
    _text(slide, f"Financial Research Automation  |  Slide {n}", Inches(0.5), Inches(7.05),
          Inches(12), Inches(0.3), 10, False, GREY)


def _table(slide, df: pd.DataFrame, left, top, width, height, size=12):
    rows, cols = df.shape[0] + 1, df.shape[1]
    tbl = slide.shapes.add_table(rows, cols, left, top, width, height).table
    for j, col in enumerate(df.columns):
        c = tbl.cell(0, j)
        c.text = str(col)
        c.fill.solid()
        c.fill.fore_color.rgb = NAVY
        para = c.text_frame.paragraphs[0]
        para.font.size, para.font.bold, para.font.color.rgb = Pt(size), True, WHITE
        para.alignment = PP_ALIGN.CENTER
    for i, (_, r) in enumerate(df.iterrows(), start=1):
        for j, v in enumerate(r):
            c = tbl.cell(i, j)
            c.text = str(v)
            c.fill.solid()
            c.fill.fore_color.rgb = LIGHT if i % 2 else WHITE
            para = c.text_frame.paragraphs[0]
            para.font.size, para.font.color.rgb = Pt(size), GREY
            para.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER


def _fmt(v, kind="num"):
    if pd.isna(v):
        return "-"
    return {"num": f"{v:,.0f}", "pct": f"{v:.1f}%", "dec": f"{v:.2f}"}[kind]


def build_report(wide: pd.DataFrame, snap: pd.DataFrame, charts: dict[str, Path],
                 source_label: str = "Screener.in") -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH
    n = 0

    # 1. Title
    s = _blank(prs); n += 1
    bg = s.shapes.add_shape(1, 0, 0, SW, SH)
    bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
    acc = s.shapes.add_shape(1, Inches(0.8), Inches(3.55), Inches(2.5), Inches(0.08))
    acc.fill.solid(); acc.fill.fore_color.rgb = ORANGE; acc.line.fill.background()
    _text(s, "IT Sector Quarterly Financial Report", Inches(0.8), Inches(2.2), Inches(11.5),
          Inches(1.3), 40, True, WHITE)
    comps = ", ".join(sorted(wide["Company"].unique()))
    _text(s, comps, Inches(0.8), Inches(3.8), Inches(11.5), Inches(0.8), 18, False, WHITE)
    _text(s, f"Generated {datetime.now():%d %b %Y}  |  Source: {source_label}",
          Inches(0.8), Inches(6.4), Inches(11.5), Inches(0.5), 13, False, WHITE)

    # 2. Key takeaways
    s = _blank(prs); n += 1
    _header(s, "Key Takeaways", "Sector-level findings from the latest quarter")
    _bullets(s, analytics.sector_summary(wide, snap), Inches(0.7), Inches(1.8), Inches(12), Inches(5), 20)
    _footer(s, n)

    # 3. Snapshot table
    s = _blank(prs); n += 1
    _header(s, "Latest Quarter Snapshot", "Rs Crore unless stated")
    t = pd.DataFrame({
        "Company": snap["Company"],
        "Quarter": snap["Quarter Label"],
        "Revenue": snap["Revenue"].map(_fmt),
        "Net Profit": snap["Net Profit"].map(_fmt),
        "OPM %": snap["OPM %"].map(lambda v: _fmt(v, "pct")),
        "Rev YoY": snap["Revenue YoY %"].map(lambda v: _fmt(v, "pct")),
        "EPS (Rs)": snap["EPS"].map(lambda v: _fmt(v, "dec")),
    })
    _table(s, t, Inches(0.5), Inches(1.7), Inches(12.3), Inches(0.5 * (len(t) + 1)), 13)
    _footer(s, n)

    # 4-8. Comparison charts
    for key, title in [("Revenue trend", "Revenue Trend"),
                       ("Margin trend", "Operating Margin Trend"),
                       ("Net profit (latest)", "Net Profit Comparison"),
                       ("Growth comparison", "Growth Comparison"),
                       ("QoQ heatmap", "QoQ Revenue Growth Heatmap")]:
        if key not in charts:
            continue
        s = _blank(prs); n += 1
        _header(s, title)
        _picture_fit(s, charts[key], Inches(0.8), Inches(1.2), Inches(11.7), Inches(5.7))
        _footer(s, n)

    # Company slides
    for comp in sorted(wide["Company"].unique()):
        key = f"company::{comp}"
        s = _blank(prs); n += 1
        _header(s, comp, "Company deep-dive")
        if key in charts:
            _picture_fit(s, charts[key], Inches(0.4), Inches(1.7), Inches(7.6), Inches(5.0), center=False)
        _bullets(s, analytics.company_insights(wide, comp), Inches(8.2), Inches(1.7),
                 Inches(4.8), Inches(5), 14)
        _footer(s, n)

    # Methodology
    s = _blank(prs); n += 1
    _header(s, "Methodology & Notes")
    _bullets(s, [
        f"Data source: {source_label} (quarterly results, consolidated).",
        "Pipeline: Selenium/BeautifulSoup scraping -> Pandas cleaning -> NumPy analytics -> Matplotlib charts -> python-pptx report.",
        "QoQ = change vs previous quarter; YoY = change vs same quarter last year.",
        "OPM % = Operating Profit / Revenue. Figures in Rs Crore.",
        "This report is informational and not investment advice.",
    ], Inches(0.7), Inches(1.5), Inches(12), Inches(5), 15)
    _footer(s, n)

    out = config.REPORT_DIR / f"IT_Financial_Report_{datetime.now():%Y%m%d_%H%M}.pptx"
    prs.save(out)
    return out
